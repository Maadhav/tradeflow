"""Build Tradeflow.pptx from the rendered slides (deck/pdf/slide-N.png) with a video embedded.

Each slide is the full-bleed image of the HTML deck. The demo slide shows the video instead of its
screenshot, with a poster frame taken from the video, so it plays on stage from the slide itself.

Usage: deck/.venv/bin/python deck/build_pptx.py <video.mp4> [demo slide number, default 7] [out.pptx]
POSTER_AT=<seconds> picks the poster frame (default 2).
"""

import os
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches

HERE = Path(__file__).resolve().parent
SLIDES = sorted((HERE / "pdf").glob("slide-*.png"), key=lambda p: int(p.stem.split("-")[1]))

video = Path(sys.argv[1]).expanduser().resolve()
demo = int(sys.argv[2]) if len(sys.argv) > 2 else 7
out = Path(sys.argv[3]).expanduser() if len(sys.argv) > 3 else HERE / "Tradeflow.pptx"
if not video.exists():
    sys.exit(f"video not found: {video}")

poster = HERE / "pdf" / "video-poster.png"
subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-ss", os.environ.get("POSTER_AT", "2"), "-i", str(video), "-frames:v", "1", str(poster)], check=True)
probe = subprocess.run(
    ["ffprobe", "-v", "error", "-select_streams", "v:0", "-show_entries", "stream=width,height", "-of", "csv=p=0", str(video)],
    check=True, capture_output=True, text=True,
).stdout.strip().split(",")
aspect = int(probe[0]) / int(probe[1])

# The demo slide keeps its title; everything below it becomes plain background for the video.
TITLE_BOTTOM = 130  # px on the 1280 x 720 slide
demo_bg = HERE / "pdf" / "demo-background.png"
subprocess.run(
    ["ffmpeg", "-y", "-loglevel", "error", "-i", str(SLIDES[demo - 1]), "-vf",
     f"drawbox=x=0:y={TITLE_BOTTOM}*ih/720:w=iw:h=ih:color=0x14202b:t=fill", str(demo_bg)],
    check=True,
)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height

for n, image in enumerate(SLIDES, start=1):
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(str(demo_bg if n == demo else image), 0, 0, W, H)
    if n == demo:
        # As large as fits below the title, keeping the recording's own aspect ratio, centred.
        px = W / 1280
        max_w, max_h = 1136, 720 - TITLE_BOTTOM - 28
        w, h = (max_w, max_w / aspect) if max_w / aspect <= max_h else (max_h * aspect, max_h)
        left, top = Emu(int((1280 - w) / 2 * px)), Emu(int(TITLE_BOTTOM * px))
        slide.shapes.add_movie(str(video), left, top, Emu(int(w * px)), Emu(int(h * px)), poster_frame_image=str(poster), mime_type="video/mp4")

prs.save(out)
print(f"wrote {out} with {len(SLIDES)} slides, video on slide {demo}")
