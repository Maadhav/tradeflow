#!/usr/bin/env python3
"""Build the Tradeflow pitch deck (deck/Tradeflow.pptx, 16:9, 13.333 x 7.5 in).

Usage (from the repo root):
    python3 -m venv deck/.venv && deck/.venv/bin/pip install python-pptx pillow
    deck/.venv/bin/python deck/build_deck.py

Needs ffmpeg on PATH: product stills, the two demo videos and their posters are made from the demo
recordings in demo/out/ and cached in deck/assets/. Slide 4 embeds a short cut of the 2:46 demo
(see CUT) that starts automatically when the slide opens; the appendix slide embeds the whole 2:46.
In both, the last sentence of the closing caption is painted out (MASK), so GSOS is named only on
the closing slide. soffice and pdftoppm, when present, render slide 1 as the package thumbnail.

Every number on the slides is checked against the repo: evidence/README.md (14 workflow runs,
14 Sepolia transactions, APR 9% in run 1 and 19% in run 12), contracts/test (8 tests),
scripts/e2e.ts (12 steps). The one market figure is the ADB 2023 Trade Finance Gaps survey, cited
on the slide.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import io
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Inches, Pt

DECK = Path(__file__).resolve().parent
REPO = DECK.parent
ASSETS = DECK / "assets"
SHORT = REPO / "demo/out/tradeflow-short.mp4"  # 2:46, H.264, 1440x900
FULL = REPO / "demo/out/final-sepolia.mp4"  # the full Sepolia take, source of the stills
OUT = DECK / "Tradeflow.pptx"

# ---------- product identity (services/rails/web/styles.css) ----------
NAVY = RGBColor(0x14, 0x20, 0x2B)
CRANE = RGBColor(0xF2, 0xB7, 0x05)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xF2, 0xF4, 0xF6)
INK = RGBColor(0x3A, 0x4A, 0x59)
MUTED = RGBColor(0x5B, 0x6A, 0x78)  # captions on paper (5:1 contrast)
RULE = RGBColor(0xB9, 0xC3, 0xCC)
CARD_LINE = RGBColor(0xD9, 0xDF, 0xE5)
NAVY_2 = RGBColor(0x24, 0x34, 0x44)  # a raised surface on navy
MIST = RGBColor(0xA9, 0xB6, 0xC2)  # secondary text on navy (8:1 contrast)
FONT = "Arial"

SLIDE_W, SLIDE_H = 13.333, 7.5
M = 0.75  # side margin
CONTENT_W = SLIDE_W - 2 * M

# The pitch cut of the 2:46 demo, as (start, end) seconds in demo/out/tradeflow-short.mp4, about
# 1:27 in all: the invoice and the buyer's confirmation, the listing run with its TEE terminal,
# a bank deposit checked and credited, full funding and payout, two-source repayment, the books
# mismatch that pauses funding onchain, then the late loan that freezes the business, ending on the
# marketplace. The sign-up typing, the second invoice and the wallet lender are left for the appendix.
CUT = [
    (21.0, 29.4),  # the invoice is entered; nothing lists until the buyer confirms
    (30.4, 34.8),  # the buyer opens the link and confirms
    (35.0, 59.6),  # listing workflow live, confidential credit review, listed, then the CLI terminal
    (72.6, 80.6),  # Ana's bank deposit is checked against the Data Feed; she holds loan notes
    (101.2, 106.2),  # fully funded: the log trigger pays the business
    (116.0, 123.8),  # repayment confirmed by two sources
    (126.0, 135.8),  # the books disagree: funding paused onchain, and the transaction on Sepolia
    (139.0, 157.8),  # the buyer does not pay: late, frozen; the marketplace at the end
]
# The closing caption (from 154.26 s) ends "...writes through the forwarder. Built by CodeDecoders,
# the team behind GSOS." Paint out the second sentence with the caption band's colour.
MASK = "drawbox=x=211:y=862:w=326:h=20:color=0x14202C:t=fill:enable='gte(t,154.26)'"


# ---------- assets ----------
def frame(video: Path, t: float, out: Path) -> Path:
    if not out.exists():
        subprocess.run(
            ["ffmpeg", "-v", "error", "-y", "-ss", str(t), "-i", str(video), "-frames:v", "1", str(out)],
            check=True,
        )
    return out


def duration(video: Path) -> float:
    r = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0",
                        str(video)], check=True, capture_output=True, text=True)
    return float(r.stdout.strip())


def encode(filters: str, out: Path) -> Path:
    if not out.exists():
        subprocess.run(
            ["ffmpeg", "-v", "error", "-y", "-i", str(SHORT), "-filter_complex", filters, "-map", "[out]",
             "-c:v", "libx264", "-preset", "slow", "-crf", "20", "-pix_fmt", "yuv420p", "-r", "25",
             "-movflags", "+faststart", "-an", str(out)],
            check=True,
        )
    return out


def build_videos() -> dict[str, Path]:
    key = hashlib.sha1(repr((CUT, MASK)).encode()).hexdigest()[:8]
    n = len(CUT)
    parts = "".join(f"[s{i}]trim=start={a}:end={b},setpts=PTS-STARTPTS[v{i}];" for i, (a, b) in enumerate(CUT))
    cut = encode(
        f"[0:v]{MASK},split={n}" + "".join(f"[s{i}]" for i in range(n)) + ";" + parts
        + "".join(f"[v{i}]" for i in range(n)) + f"concat=n={n}:v=1:a=0[out]",
        ASSETS / f"demo-cut-{key}.mp4",
    )
    full = encode(f"[0:v]{MASK}[out]", ASSETS / f"demo-full-{key}.mp4")
    return {
        "cut": cut,
        "full": full,
        # posters: the marketplace at the end of the cut, and the TEE terminal for the appendix
        "cut_poster": frame(cut, duration(cut) - 1.0, ASSETS / f"poster-cut-{key}.png"),
        "full_poster": frame(full, 55.0, ASSETS / f"poster-full-{key}.png"),
    }


def build_assets() -> dict[str, Path]:
    ASSETS.mkdir(exist_ok=True)
    raw = {
        "loan": frame(FULL, 154, ASSETS / "raw-loan-154s.png"),  # loan page INV-2026-0142
        "terminal": frame(FULL, 88, ASSETS / "raw-terminal-88s.png"),  # cre workflow simulate, TEE banner
        "explorer": frame(FULL, 490, ASSETS / "raw-explorer-490s.png"),  # Sepolia explorer, run 10 report
    }
    out = build_videos()

    # Loan page: below the app bar, down to the end of the verification card and the lender form's
    # hint, above the Fund button. Paint out the mouse pointer that sits in an empty table cell.
    im = Image.open(raw["loan"]).convert("RGB")
    bg = im.getpixel((380, 623))
    ImageDraw.Draw(im).rectangle((318, 604, 356, 642), fill=bg)
    out["loan"] = ASSETS / "loan.png"
    im.crop((120, 70, 1320, 757)).save(out["loan"])

    # Terminal (run 1, listing): the CLI's TEE banner, the handler's step logs, the report delivered
    # to Sepolia (tx 0x22db61cd...) and the first line of the simulation result.
    out["terminal"] = ASSETS / "terminal-tee.png"
    Image.open(raw["terminal"]).convert("RGB").crop((14, 318, 832, 650)).save(out["terminal"])

    # Sepolia explorer: the monitor's signed report (run 10). The explorer's transient "block is
    # being re-synced" notice (rows 449 to 496) is cut out; nothing about the transaction is.
    out["explorer"] = ASSETS / "explorer.png"
    ex = Image.open(raw["explorer"]).convert("RGB")
    top, bottom = ex.crop((16, 222, 952, 449)), ex.crop((16, 497, 952, 812))
    spliced = Image.new("RGB", (936, top.height + bottom.height), "white")
    spliced.paste(top, (0, 0))
    spliced.paste(bottom, (0, top.height))
    spliced.save(out["explorer"])
    return out


def img_size(path: Path) -> tuple[int, int]:
    with Image.open(path) as im:
        return im.size


def mmss(seconds: float) -> str:
    s = int(round(seconds))
    return f"{s // 60}:{s % 60:02d}"


# ---------- drawing helpers ----------
def set_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def text(slide, x, y, w, h, paras, *, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None, gap=0, name=None, font=FONT):
    """A text box with no padding. `paras` is a string, or a list of paragraphs where each
    paragraph is a string or a list of (text, options) runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.line_spacing = spacing
        if gap and i < len(paras) - 1:
            p.space_after = Pt(gap)
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.color.rgb = o.get("color", color)
    return tb


def unstyle(shape):
    """Drop the theme style reference python-pptx adds to autoshapes and connectors. Fill and line
    are always set explicitly here, and the style's effect reference draws a drop shadow in some
    renderers even when the shape carries an empty effect list."""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)
    return shape


def box(slide, x, y, w, h, fill, *, line=None, radius=0.12, name=None):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is not None:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    if radius:
        s.adjustments[0] = radius / min(w, h)
    return unstyle(s)


def dot(slide, cx, cy, d, *, fill=None, ring=None, weight=2.25):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    if fill is not None:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    else:
        s.fill.background()
    if ring is not None:
        s.line.color.rgb = ring
        s.line.width = Pt(weight)
    else:
        s.line.fill.background()
    return unstyle(s)


def dashed(slide, x1, y1, x2, y2, color, weight=2.25):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(weight)
    c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return unstyle(c)


def lane(slide, x1, x2, y, *, d=0.3, line=NAVY, ring=NAVY, weight=2.25, stops=(), stop_fill=None):
    """The trade lane from the product's wordmark: a crane-yellow origin dot, a dashed route and a
    hollow destination ring. Optional intermediate stops (x centres) sit on the route."""
    pad = d / 2 + 0.07
    dashed(slide, x1 + pad, y, x2 - pad, y, line, weight)
    for sx in stops:
        dot(slide, sx, y, d * 0.62, fill=stop_fill or line)
    dot(slide, x1, y, d, fill=CRANE)
    dot(slide, x2, y, d, ring=ring, weight=weight * 1.15)


def vlane(slide, x, y1, y2, *, d=0.22, line=NAVY, ring=NAVY, weight=2.0, stops=(), stop_fill=None):
    pad = d / 2 + 0.06
    dashed(slide, x, y1 + pad, x, y2 - pad, line, weight)
    for sy in stops:
        dot(slide, x, sy, d * 0.62, fill=stop_fill or line)
    dot(slide, x, y1, d, fill=CRANE)
    dot(slide, x, y2, d, ring=ring, weight=weight * 1.15)


def picture(slide, path, x, y, w, *, border=None, name=None):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    if name:
        pic.name = name
    if border is not None:
        pic.line.color.rgb = border
        pic.line.width = Pt(0.75)
    return pic


def title(slide, words, *, color=NAVY):
    return text(slide, M, 0.62, CONTENT_W, 0.8, words, size=40, bold=True, color=color, name="Title")


def notes(slide, words: str) -> None:
    slide.notes_slide.notes_text_frame.text = words


def blank(prs, bg):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, bg)
    return s


# ---------- autoplay for the embedded video ----------
AUTOPLAY_TIMING = """
<p:timing {ns}>
  <p:tnLst>
    <p:par>
      <p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
        <p:childTnLst>
          <p:seq concurrent="1" nextAc="seek">
            <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
              <p:childTnLst>
                <p:par>
                  <p:cTn id="3" fill="hold">
                    <p:stCondLst>
                      <p:cond delay="indefinite"/>
                      <p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond>
                    </p:stCondLst>
                    <p:childTnLst>
                      <p:par>
                        <p:cTn id="4" fill="hold">
                          <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                          <p:childTnLst>
                            <p:par>
                              <p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect">
                                <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                                <p:childTnLst>
                                  <p:cmd type="call" cmd="playFrom(0.0)">
                                    <p:cBhvr>
                                      <p:cTn id="6" dur="{dur}" fill="hold"/>
                                      <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
                                    </p:cBhvr>
                                  </p:cmd>
                                </p:childTnLst>
                              </p:cTn>
                            </p:par>
                          </p:childTnLst>
                        </p:cTn>
                      </p:par>
                    </p:childTnLst>
                  </p:cTn>
                </p:par>
              </p:childTnLst>
            </p:cTn>
            <p:prevCondLst>
              <p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>
            </p:prevCondLst>
            <p:nextCondLst>
              <p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>
            </p:nextCondLst>
          </p:seq>
          <p:video>
            <p:cMediaNode vol="80000">
              <p:cTn id="7" fill="hold" display="0">
                <p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>
                <p:endCondLst>
                  <p:cond evt="onStopAudio" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>
                </p:endCondLst>
              </p:cTn>
              <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
            </p:cMediaNode>
          </p:video>
        </p:childTnLst>
      </p:cTn>
    </p:par>
  </p:tnLst>
</p:timing>
"""


def autoplay(slide, movie, seconds: float) -> None:
    """Replace python-pptx's click-to-play timing with PowerPoint's 'Start: Automatically'."""
    sld = slide._element
    old = sld.find(qn("p:timing"))
    new = parse_xml(AUTOPLAY_TIMING.format(ns=nsdecls("p"), spid=movie.shape_id, dur=int(seconds * 1000)))
    old.addprevious(new)
    sld.remove(old)


def video_column(s, a, video, poster, *, heading, length, sub, chapters=(), auto=True, heading_size=32):
    """A demo video framed on navy at the left, its heading and chapter rail on the right."""
    w, h = 9.36, 5.85
    x, y = M, (SLIDE_H - h) / 2
    box(s, x - 0.08, y - 0.08, w + 0.16, h + 0.16, NAVY_2, radius=0.08, name="Screen")
    movie = s.shapes.add_movie(str(video), Inches(x), Inches(y), Inches(w), Inches(h),
                               poster_frame_image=str(poster), mime_type="video/mp4")
    movie.name = heading
    if auto:
        autoplay(s, movie, duration(video))
    cx = x + w + 0.45
    cw = SLIDE_W - M - cx
    text(s, cx, y - 0.05, cw, 0.6, heading, size=heading_size, bold=True, color=WHITE)
    text(s, cx, y + 0.62, cw, 0.4, length, size=20, color=CRANE, bold=True)
    if sub:
        text(s, cx, y + 1.05, cw, 0.6, sub, size=12, color=MIST, spacing=1.1)
    if chapters:
        y1, y2 = y + 1.85, y + h - 0.2
        step = (y2 - y1) / (len(chapters) - 1)
        ys = [y1 + i * step for i in range(len(chapters))]
        vlane(s, cx + 0.12, y1, y2, d=0.24, line=MIST, ring=WHITE, stops=ys[1:-1], stop_fill=WHITE)
        for cy, label in zip(ys, chapters):
            text(s, cx + 0.45, cy - 0.2, cw - 0.45, 0.4, label, size=17, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    return movie


# ---------- slides ----------
def slide_title(prs, a):
    s = blank(prs, NAVY)
    lane(s, M + 0.17, 5.4, 1.95, d=0.34, line=WHITE, ring=WHITE)
    text(s, M, 2.35, 6.2, 1.35, "Tradeflow", size=80, bold=True, color=WHITE, name="Wordmark")
    text(s, M, 3.75, 6.0, 1.3, ["Fund the goods", "already on their way."], size=32, bold=True,
         color=CRANE, spacing=1.0, name="Tagline")
    text(s, M, 5.75, 6.0, 0.4, "Maadhav Sharma, Founder & CEO, CodeDecoders", size=16, color=WHITE)
    text(s, M, 6.18, 6.0, 0.4, "TOKEN2049 Origins, Singapore. Built on Chainlink CRE.", size=14, color=MIST)
    pw = SLIDE_W - M - 6.95
    iw, ih = img_size(a["loan"])
    picture(s, a["loan"], 6.95, 1.78, pw, name="Loan page")
    text(s, 6.95, 1.78 + pw * ih / iw + 0.15, pw, 0.3,
         "A listed loan in the app: verified, priced and open for funding", size=12, color=MIST)
    notes(s, "Tradeflow funds goods that are already on their way. A business gets an advance on a "
             "confirmed invoice, bill of lading or equipment order, lenders fund it, and every "
             "decision that depends on off-chain facts is made by a Chainlink CRE workflow.")


def slide_problem(prs, a):
    s = blank(prs, PAPER)
    title(s, "Shipped today, paid months later")
    y = 2.55
    lane(s, M + 0.2, SLIDE_W - M - 0.2, y, d=0.4, line=INK, ring=NAVY, weight=2.5)
    text(s, M, y + 0.38, 4.0, 0.4, "Goods shipped", size=18, bold=True, color=NAVY)
    text(s, SLIDE_W - M - 4.0, y + 0.38, 4.0, 0.4, "Buyer pays", size=18, bold=True, color=NAVY,
         align=PP_ALIGN.RIGHT)
    text(s, M, y + 1.05, CONTENT_W, 0.45,
         "In between, the exporter still has to pay for the next order.", size=22, color=INK)
    text(s, M, 4.65, 7.0, 1.0, "$2.5 trillion", size=60, bold=True, color=NAVY, name="Gap")
    text(s, M, 5.65, 7.0, 0.45, "Global trade finance gap, 2022", size=22, color=INK)
    text(s, M, 6.55, CONTENT_W, 0.3,
         "Source: Asian Development Bank, 2023 Trade Finance Gaps, Growth, and Jobs Survey.",
         size=12, color=MUTED)
    notes(s, "Exporters, shippers and manufacturers wait months to be paid for goods they have "
             "already shipped, and still have to pay for the next order in between. The Asian "
             "Development Bank puts the global trade finance gap at 2.5 trillion dollars in 2022.")


def slide_flow(prs, a):
    s = blank(prs, PAPER)
    title(s, "Sell the wait to lenders")
    steps = [
        ("Business", "Asks for an advance"),
        ("Buyer", "Confirms by link"),
        ("Lenders", "Fund in fiat or USDC"),
        ("Business", "Is paid out in fiat"),
        ("Buyer", "Pays, lenders repaid"),
    ]
    col_w, gap = 2.21, 0.195
    xs = [M + i * (col_w + gap) for i in range(len(steps))]
    y = 2.1
    d = 0.34
    lane(s, xs[0] + d / 2, xs[-1] + d / 2, y, d=d, line=INK, ring=NAVY, weight=2.25,
         stops=[x + d / 2 for x in xs[1:-1]], stop_fill=NAVY)
    for x, (who, what) in zip(xs, steps):
        text(s, x, y + 0.42, col_w, 0.5, who, size=24, bold=True, color=NAVY)
        text(s, x, y + 0.95, col_w + gap - 0.03, 0.4, what, size=17, color=INK)
    # who benefits
    by = 4.0
    text(s, M, by, CONTENT_W, 0.35, "Who benefits", size=16, bold=True, color=MUTED)
    benefits = [
        ("Exporters", "Paid when the goods ship"),
        ("Lenders", "Short, document-backed yield: 9% to 19% APR in the demo"),
        ("Tradeflow", "A fee on each advance"),
    ]
    bgap = 0.3
    bw = (CONTENT_W - 2 * bgap) / 3
    for i, (who, what) in enumerate(benefits):
        bx = M + i * (bw + bgap)
        box(s, bx, by + 0.42, bw, 1.3, WHITE, line=CARD_LINE, radius=0.1, name=f"{who} benefit")
        dot(s, bx + 0.36, by + 0.42 + 0.38, 0.16, fill=CRANE)
        text(s, bx + 0.6, by + 0.42 + 0.2, bw - 0.85, 0.4, who, size=20, bold=True, color=NAVY)
        text(s, bx + 0.6, by + 0.42 + 0.62, bw - 0.85, 0.6, what, size=15, color=INK, spacing=1.05)
    text(s, M, 6.25, CONTENT_W, 0.45, [[
        ("Every off-chain decision is a ", {}),
        ("Chainlink CRE workflow", {"bold": True, "color": NAVY}),
        (", written onchain as a signed report.", {}),
    ]], size=20, color=INK)
    notes(s, "The business enters a confirmed invoice, bill of lading or equipment order. The buyer "
             "confirms it through a link, and nothing lists until they do. Lenders fund it by bank "
             "transfer, on-ramped to USDC, or with USDC from a wallet. The business is paid out in "
             "fiat as soon as the loan is fully funded, and when the buyer pays, lenders are repaid "
             "with interest. Exporters are paid when the goods ship, not months later. Lenders get "
             "short, document-backed yield: the two demo loans priced at 9 and 19 percent APR. "
             "Tradeflow's model is a fee on each advance.")


def slide_demo(prs, a):
    s = blank(prs, NAVY)
    video_column(s, a, a["cut"], a["cut_poster"], heading="Demo", length=mmss(duration(a["cut"])),
                 sub=f"Full {mmss(duration(a['full']))} cut: last slide",
                 chapters=["Listed", "Funded", "Paid out", "Repaid", "Paused", "Late, frozen"])
    notes(s, "The video starts on its own and runs on Ethereum Sepolia. The buyer confirms the "
             "invoice and the listing workflow runs: the credit review is a TEE-declared handler, and "
             "the terminal shows the same run in the CRE CLI. A bank deposit is checked against the "
             "Chainlink Data Feed, the loan funds and the business is paid out. Repayment is "
             "confirmed by two independent sources. Then the bank books a deposit that never reached "
             "the chain, and the monitor pauses funding onchain. A second loan goes late and the "
             "business is frozen. The longer cut is on the last slide.")


def slide_workflows(prs, a):
    s = blank(prs, PAPER)
    title(s, "CRE makes every off-chain decision")
    cards = [
        ("listing", ["HTTP trigger", "TEE handler", "Data Feed: EUR/USD"],
         "Is the document real, confirmed and creditworthy? Then list the loan."),
        ("lender", ["HTTP trigger", "Confidential HTTP", "Data Feed, 1% FX check"],
         "Is the lender verified? Did the bank deposit really arrive?"),
        ("settlement", ["Log + HTTP triggers", "Two-source consensus"],
         "Pay the business once funded. Repaid only when two sources agree."),
        ("monitor", ["Cron trigger", "Circuit breaker"],
         "Late loan? Books out of balance? Freeze the business or pause funding."),
    ]
    gap = 0.3
    cw = (CONTENT_W - gap * 3) / 4
    xs = [M + i * (cw + gap) for i in range(4)]
    ly, cy, ch = 1.95, 2.4, 3.55
    d = 0.3
    pad = 0.3
    lane(s, xs[0] + pad + d / 2, xs[-1] + pad + d / 2, ly, d=d, line=INK, ring=NAVY,
         stops=[x + pad + d / 2 for x in xs[1:-1]], stop_fill=NAVY)
    for x, (name, trig, decides) in zip(xs, cards):
        box(s, x, cy, cw, ch, WHITE, line=CARD_LINE, radius=0.1, name=f"{name} card")
        text(s, x + pad, cy + 0.3, cw - 2 * pad, 0.5, name, size=26, bold=True, color=NAVY)
        text(s, x + pad, cy + 0.9, cw - 2 * pad, 0.8, trig, size=13, color=MUTED)
        text(s, x + pad, cy + 1.85, cw - 2 * pad, 1.6, decides, size=17, color=INK, spacing=1.1)
    text(s, M, 6.25, CONTENT_W, 0.45,
         "Each run writes a signed report onchain through the Chainlink CRE forwarder.",
         size=18, color=INK)
    notes(s, "Four workflows, seven handlers. Listing verifies and prices the document, and prices "
             "euro documents with the Chainlink EUR/USD Data Feed. Lender runs KYC over Confidential "
             "HTTP and only credits a bank deposit when the on-ramp rate is within 1 percent of the "
             "feed. Settlement pays out on the funded log event, and an HTTP-triggered run confirms "
             "repayment from two independent sources. Monitor runs on a cron, catches late loans and "
             "reconciles the books, with an onchain circuit breaker.")


def slide_confidential(prs, a):
    s = blank(prs, PAPER)
    title(s, "Confidential by design")
    lw, rx = 6.0, 8.15
    rw = SLIDE_W - M - rx
    top, bh = 1.65, 2.2
    # what stays in the enclave
    box(s, M, top, lw, bh, NAVY, radius=0.12, name="Enclave")
    text(s, M + 0.32, top + 0.25, lw - 0.6, 0.45, "Stays in the enclave", size=22, bold=True, color=WHITE)
    text(s, M + 0.32, top + 0.7, lw - 0.6, 0.35, "AWS Nitro on a deployed DON, declared with handlerInTee",
         size=14, color=MIST)
    for i, item in enumerate(["Revenue and open debts", "Payment history", "Sanctions screen"]):
        iy = top + 1.17 + i * 0.32
        dot(s, M + 0.4, iy + 0.14, 0.12, fill=CRANE)
        text(s, M + 0.6, iy, lw - 0.9, 0.3, item, size=17, color=WHITE)
    # only the verdict crosses
    lane(s, M + lw + 0.3, rx - 0.3, top + bh / 2, d=0.26, line=INK, ring=NAVY)
    box(s, rx, top, rw, bh, WHITE, line=CARD_LINE, radius=0.12, name="Leaves")
    text(s, rx + 0.32, top + 0.25, rw - 0.6, 0.45, "Only the verdict leaves", size=22, bold=True, color=NAVY)
    text(s, rx + 0.32, top + 0.7, rw - 0.6, 0.35, "Written onchain with the document's terms", size=14,
         color=MUTED)
    for i, item in enumerate(["Grade", "APR", "Advance"]):
        iy = top + 1.17 + i * 0.32
        dot(s, rx + 0.4, iy + 0.14, 0.13, ring=NAVY, weight=1.75)
        text(s, rx + 0.6, iy, rw - 0.9, 0.3, item, size=17, color=INK)
    # evidence and Confidential HTTP
    by = top + bh + 0.3
    iw, ih = img_size(a["terminal"])
    th = lw * ih / iw
    picture(s, a["terminal"], M, by, lw, name="Terminal")
    text(s, M, by + th + 0.1, lw, 0.45,
         "CRE CLI simulation of run 1: the handler requests AWS Nitro, the simulator runs it locally "
         "and says so, then delivers the signed report to Sepolia.", size=12, color=MUTED, spacing=1.05)
    box(s, rx, by, rw, th, WHITE, line=CARD_LINE, radius=0.12, name="Confidential HTTP")
    text(s, rx + 0.32, by + 0.25, rw - 0.6, 0.45, "Confidential HTTP", size=22, bold=True, color=NAVY)
    text(s, rx + 0.32, by + 0.78, rw - 0.6, th - 0.95,
         "KYC and bank deposit checks call their APIs with a Vault DON secret. On a deployed DON it "
         "is injected inside the enclave, never held in workflow code.", size=16, color=INK, spacing=1.1)
    notes(s, "The listing handler is declared for TEE execution with handlerInTee. On a deployed DON "
             "the credit file, the registry response and the sanctions result are processed inside an "
             "AWS Nitro enclave, so node operators never see a business's revenue, debts or payment "
             "history. From the credit file only the grade, APR and advance come out, written onchain "
             "with the document's own terms and hash. Lender KYC and bank deposits use Confidential "
             "HTTP with a Vault DON secret. What you see here is the CRE CLI simulation: it runs the "
             "same handler locally, says in its banner that it is not a real TEE, and then delivers "
             "the report to Sepolia. Deploy access is not enabled for the hackathon.")


def slide_onchain(prs, a):
    s = blank(prs, PAPER)
    title(s, "Lenders do not take our word for it")
    items = [
        ("KYC-gated loan notes", "ERC-1155 notes that only move between verified holders.",
         "contracts/src/LoanNotes.sol"),
        ("Enforced by code", "A late loan freezes the business. A books mismatch pauses all new funding.",
         "frozenBorrower, fundingPaused in TradeflowMarket"),
        ("No double financing", "The document hash is recorded onchain. The same invoice cannot list twice.",
         "usedRef in TradeflowMarket, checked by the listing workflow"),
        ("Every decision is signed", "A CRE report through the forwarder. The market rejects anything "
         "outside policy.", "TradeflowMarket is a CRE ReceiverTemplate"),
    ]
    gap = 0.3
    cw = (CONTENT_W - gap) / 2
    ch = 2.2
    for i, (lead, body, ref) in enumerate(items):
        x = M + (i % 2) * (cw + gap)
        y = 1.8 + (i // 2) * (ch + gap)
        box(s, x, y, cw, ch, WHITE, line=CARD_LINE, radius=0.12, name=lead)
        dot(s, x + 0.47, y + 0.56, 0.2, fill=CRANE)
        text(s, x + 0.8, y + 0.36, cw - 1.1, 0.45, lead, size=24, bold=True, color=NAVY)
        text(s, x + 0.8, y + 0.95, cw - 1.1, 0.75, body, size=18, color=INK, spacing=1.1)
        text(s, x + 0.8, y + ch - 0.47, cw - 1.1, 0.25, ref, size=12, color=MUTED)
    notes(s, "The chain is where the rules are enforced, not just recorded. Loan notes can only be "
             "held by verified people. A late loan freezes the business onchain, and a reconciliation "
             "mismatch trips a circuit breaker. Every financed document's hash is recorded, so the "
             "same invoice cannot be financed twice. Every state change arrives as a signed CRE "
             "report, and the contract checks it against policy. Unlike onchain credit pools where a "
             "trusted originator attests to its loans, each Tradeflow loan is verified by a CRE "
             "workflow before it lists.")


def slide_proof(prs, a):
    s = blank(prs, PAPER)
    title(s, "It runs on Ethereum Sepolia")
    lw = 4.85
    text(s, M, 1.7, lw, 0.85, "14", size=54, bold=True, color=NAVY)
    text(s, M, 2.55, lw, 0.35, "CRE workflow runs, 14 Sepolia transactions", size=16, color=INK)
    text(s, M, 2.92, lw, 0.3, "cre workflow simulate --broadcast", size=14, bold=True, color=NAVY,
         font="Courier New")
    text(s, M, 3.4, lw, 0.85, "8 of 8", size=54, bold=True, color=NAVY)
    text(s, M, 4.25, lw, 0.35, "contract tests pass (forge test)", size=16, color=INK)
    text(s, M, 4.75, lw, 0.85, "12", size=54, bold=True, color=NAVY)
    text(s, M, 5.6, lw, 0.35, "step end-to-end test, on a local Sepolia fork", size=16, color=INK)
    text(s, M, 6.15, lw, 0.4, "github.com/Maadhav/tradeflow", size=20, bold=True, color=NAVY)
    text(s, M, 6.57, lw, 0.3, "Every run's log and transaction hash: evidence/", size=14, color=MUTED)
    rx = 5.9
    pw = SLIDE_W - M - rx
    iw, ih = img_size(a["explorer"])
    ph = pw * ih / iw
    py = 1.8
    picture(s, a["explorer"], rx, py, pw, border=CARD_LINE, name="Sepolia explorer")
    text(s, rx, py + ph + 0.12, pw, 0.45,
         "Monitor run 10 on Sepolia: the books disagree, so a signed CRE report pauses funding "
         "onchain. The CLI simulator delivers through its Sepolia MockKeystoneForwarder.",
         size=12, color=MUTED, spacing=1.05)
    text(s, rx, 6.4, pw, 0.5,
         "Banks, KYC, credit bureau and on-ramp are simulated services. The workflows, contracts and "
         "Sepolia transactions are real. USDC is a test token.", size=12, bold=True, color=INK,
         spacing=1.05)
    notes(s, "Everything was run with the CRE CLI, cre workflow simulate with broadcast, which writes "
             "real Sepolia transactions: 14 workflow runs across the four workflows. One run was a "
             "rejection with no write, and one monitor run wrote two, so 14 transactions. Deploy "
             "access is not enabled, so reports go through the simulator's Sepolia forwarder. 8 "
             "contract tests pass, and a 12 step end-to-end test drives the whole lifecycle on a local "
             "fork of Sepolia. The banks, KYC provider, credit bureau and on-ramp are simulated "
             "services; the stablecoin is a test token. Logs and transaction hashes are in the "
             "evidence folder of the repo.")


def slide_next(prs, a):
    s = blank(prs, NAVY)
    title(s, "What is next", color=WHITE)
    steps = [
        ("Deploy on CRE", "On the CRE network, with confidential compute for the listing handler"),
        ("Banking partners", "Real banks, on-ramps and payout providers behind the same workflows"),
        ("More documents", "More document types and trade corridors, the same CRE decisions"),
    ]
    gap = 0.4
    cw = (CONTENT_W - gap * 2) / 3
    xs = [M + i * (cw + gap) for i in range(3)]
    y = 2.0
    d = 0.32
    lane(s, xs[0] + d / 2, xs[-1] + d / 2, y, d=d, line=MIST, ring=WHITE, stops=[xs[1] + d / 2],
         stop_fill=WHITE)
    for x, (head, body) in zip(xs, steps):
        text(s, x, y + 0.45, cw, 0.5, head, size=22, bold=True, color=WHITE)
        text(s, x, y + 1.0, cw - 0.2, 0.9, body, size=16, color=MIST, spacing=1.1)
    # closing
    cy = 4.85
    dot(s, M + 0.13, cy + 0.42, 0.22, fill=CRANE)
    dashed(s, M + 0.32, cy + 0.42, M + 0.62, cy + 0.42, WHITE, 2.0)
    dot(s, M + 0.8, cy + 0.42, 0.22, ring=WHITE, weight=2.25)
    text(s, M + 1.1, cy, 5.0, 0.8, "Tradeflow", size=40, bold=True, color=WHITE)
    text(s, M, cy + 0.95, 6.0, 0.5, "Fund the goods already on their way.", size=22, bold=True, color=CRANE)
    rx = 7.6
    rw = SLIDE_W - M - rx
    text(s, rx, cy + 0.05, rw, 0.4, "Maadhav Sharma", size=20, bold=True, color=WHITE)
    text(s, rx, cy + 0.47, rw, 0.35, "Founder & CEO, CodeDecoders", size=16, color=MIST)
    text(s, rx, cy + 1.0, rw, 0.35, "Built by CodeDecoders, the team behind GSOS", size=16, color=WHITE)
    text(s, rx, cy + 1.38, rw, 0.35, "gsos.io    github.com/Maadhav/tradeflow", size=14, color=MIST)
    notes(s, "Next: deploy on the CRE network with confidential compute, plug in real banking, "
             "on-ramp and payout partners, and add more document types and corridors. Built by "
             "CodeDecoders, the team behind GSOS. Thank you.")


def slide_appendix(prs, a):
    s = blank(prs, NAVY)
    video_column(s, a, a["full"], a["full_poster"], heading="Appendix", heading_size=28,
                 length=mmss(duration(a["full"])),
                 sub="Full walkthrough: the longer cut of the same Sepolia take. Sign-up, the buyer "
                     "link, a bank lender and a wallet lender, payout, repayment, the circuit breaker "
                     "and the late loan. Click the video to play.", auto=False)
    notes(s, "Appendix for reviewers, not part of the live pitch: the longer cut of the same Sepolia "
             "take shown on the demo slide.")


# ---------- build ----------
DASHES = re.compile("[‒–—―]")
FORBIDDEN = re.compile(r"60 to 180", re.I)
CLOSING_SLIDE = 9  # the only slide that may name GSOS


def check_text(prs) -> None:
    for n, slide in enumerate(prs.slides, 1):
        texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
        texts.append(slide.notes_slide.notes_text_frame.text)
        for t in texts:
            if DASHES.search(t):
                sys.exit(f"slide {n}: en or em dash in {t!r}")
            if FORBIDDEN.search(t):
                sys.exit(f"slide {n}: forbidden text in {t!r}")
            if "GSOS" in t and n != CLOSING_SLIDE:
                sys.exit(f"slide {n}: GSOS outside the closing slide in {t!r}")


def thumbnail(path: Path) -> bytes | None:
    """Slide 1 as a 256x144 JPEG, rendered with LibreOffice; None when the tools are missing."""
    if not (shutil.which("soffice") and shutil.which("pdftoppm")):
        print("soffice or pdftoppm not found: keeping the blank package thumbnail")
        return None
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, str(path)],
                       check=True, capture_output=True)
        pdf = Path(tmp) / (path.stem + ".pdf")
        subprocess.run(["pdftoppm", "-png", "-r", "40", "-f", "1", "-l", "1", str(pdf), f"{tmp}/t"],
                       check=True)
        png = next(Path(tmp).glob("t*.png"))
        im = Image.open(png).convert("RGB")
        im.thumbnail((256, 256))
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=88)
        return buf.getvalue()


def patch_package(path: Path, slides: int, videos: list[Path]) -> None:
    """Theme fonts to Arial (titles, notes and anything unstyled), a 16:9 slide size and app
    properties, the slide 1 thumbnail, then confirm both videos are embedded."""
    thumb = thumbnail(path)
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.startswith("ppt/theme/") and item.filename.endswith(".xml"):
                data = re.sub(rb'typeface="Calibri[^"]*"', b'typeface="Arial"', data)
            if item.filename.startswith("ppt/slideMasters/") and item.filename.endswith(".xml"):
                # the stock template bullets its second and fourth levels with an en dash
                data = data.replace('char="–"'.encode(), 'char="•"'.encode())
            if item.filename == "ppt/presentation.xml":
                data = data.replace(b' type="screen4x3"', b"")
            if item.filename == "docProps/app.xml":
                data = data.replace(b"On-screen Show (4:3)", b"Widescreen")
                for tag, value in (("Slides", slides), ("Notes", slides), ("MMClips", len(videos))):
                    data = re.sub(rf"<{tag}>\d+</{tag}>".encode(), f"<{tag}>{value}</{tag}>".encode(), data)
            if item.filename == "docProps/thumbnail.jpeg" and thumb:
                data = thumb
            dst.writestr(item, data)
    tmp.replace(path)
    with zipfile.ZipFile(path) as z:
        media = {i.file_size for i in z.infolist() if i.filename.startswith("ppt/media/") and i.filename.endswith(".mp4")}
        for v in videos:
            if v.stat().st_size not in media:
                sys.exit(f"{v.name} is not embedded in ppt/media")
            print(f"embedded {v.name} ({v.stat().st_size:,} bytes, {mmss(duration(v))})")


def main() -> None:
    assets = build_assets()
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    now = dt.datetime.now(dt.timezone.utc).replace(microsecond=0, tzinfo=None)
    cp = prs.core_properties
    cp.title = "Tradeflow"
    cp.author = cp.last_modified_by = "Maadhav Sharma"
    cp.subject = "Fund the goods already on their way."
    cp.comments = ""
    cp.revision = 1
    cp.created = cp.modified = now
    for build in (slide_title, slide_problem, slide_flow, slide_demo, slide_workflows,
                  slide_confidential, slide_onchain, slide_proof, slide_next, slide_appendix):
        build(prs, assets)
    check_text(prs)
    prs.save(OUT)
    patch_package(OUT, len(prs.slides), [assets["cut"], assets["full"]])
    print(f"wrote {OUT} ({OUT.stat().st_size:,} bytes, {len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
